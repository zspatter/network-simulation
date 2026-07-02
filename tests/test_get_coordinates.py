import json
import sys
from os.path import abspath, dirname, join

import openpyxl
import requests

sys.path.insert(0, join(dirname(dirname(abspath(__file__))), 'execute'))

import get_coordinates  # noqa: E402

_HEADERS = ('unique id', 'hospital name', 'city', 'state', 'region', 'latitude', 'longitude')


def _build_worksheet(rows):
    """A worksheet with exactly len(_HEADERS) header columns - no 8th column."""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(list(_HEADERS))
    for row in rows:
        worksheet.append([row[header] for header in _HEADERS])
    return worksheet


class _FakeResponse:
    def __init__(self, body=None, status_code=200):
        self.status_code = status_code
        self.text = json.dumps(body) if body is not None else ''

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.exceptions.HTTPError(f'{self.status_code} error')


def _bing_body(lat, lon):
    return {'resourceSets': [{'resources': [{'point': {'coordinates': [lat, lon]}}]}]}


def test_get_column_indices_handles_a_worksheet_with_no_eighth_column():
    # regression test: get_column_indices used to call .lower() on an empty (None)
    # header cell beyond the last real column, raising AttributeError
    worksheet = _build_worksheet([])
    columns = get_coordinates.set_default_indices()

    get_coordinates.get_column_indices(worksheet=worksheet, columns=columns)

    assert columns['city'] == 3
    assert columns['longitude'] == 7


def test_get_unique_locations_returns_distinct_city_state_region_tuples():
    rows = [
        {'unique id': 1, 'hospital name': 'A', 'city': 'Austin', 'state': 'TX', 'region': 4,
         'latitude': None, 'longitude': None},
        {'unique id': 2, 'hospital name': 'B', 'city': 'Austin', 'state': 'TX', 'region': 4,
         'latitude': None, 'longitude': None},
        {'unique id': 3, 'hospital name': 'C', 'city': 'Dallas', 'state': 'TX', 'region': 4,
         'latitude': None, 'longitude': None},
    ]
    worksheet = _build_worksheet(rows)
    columns = get_coordinates.set_default_indices()
    get_coordinates.get_column_indices(worksheet=worksheet, columns=columns)

    locations = get_coordinates.get_unique_locations(worksheet, columns)

    assert locations == {('Austin', 'TX', 4), ('Dallas', 'TX', 4)}


def test_get_coordinate_returns_coordinates_on_success(monkeypatch):
    monkeypatch.setattr(get_coordinates.requests, 'get',
                        lambda url: _FakeResponse(_bing_body(30.27, -97.74)))

    coordinates = get_coordinates.get_coordinate(city='Austin', state='TX', api_key='fake-key')

    assert coordinates == [30.27, -97.74]


def test_get_coordinate_returns_none_and_prints_on_http_error(monkeypatch, capsys):
    monkeypatch.setattr(get_coordinates.requests, 'get',
                        lambda url: _FakeResponse(status_code=404))

    coordinates = get_coordinates.get_coordinate(city='Nowhere', state='ZZ', api_key='fake-key')

    assert coordinates is None
    assert 'Nowhere' in capsys.readouterr().out


def test_get_coordinates_skips_locations_that_failed(monkeypatch):
    def fake_get(url):
        if 'Austin' in url:
            return _FakeResponse(_bing_body(30.27, -97.74))
        return _FakeResponse(status_code=404)

    monkeypatch.setattr(get_coordinates.requests, 'get', fake_get)
    locations = {('Austin', 'TX', 4), ('Nowhere', 'ZZ', 1)}

    result = get_coordinates.get_coordinates(locations, api_key='fake-key')

    assert result == {('Austin', 'TX', 4): [30.27, -97.74]}


def test_set_coordinates_writes_latitude_and_longitude_into_matching_rows():
    rows = [
        {'unique id': 1, 'hospital name': 'A', 'city': 'Austin', 'state': 'TX', 'region': 4,
         'latitude': None, 'longitude': None},
        {'unique id': 2, 'hospital name': 'B', 'city': 'Nowhere', 'state': 'ZZ', 'region': 1,
         'latitude': None, 'longitude': None},
    ]
    worksheet = _build_worksheet(rows)
    columns = get_coordinates.set_default_indices()
    get_coordinates.get_column_indices(worksheet=worksheet, columns=columns)
    locations = {('Austin', 'TX', 4): [30.27, -97.74]}

    get_coordinates.set_coordinates(worksheet, locations, columns)

    assert worksheet.cell(row=2, column=columns['latitude']).value == 30.27
    assert worksheet.cell(row=2, column=columns['longitude']).value == -97.74
    # Nowhere/ZZ had no resolved coordinates - left untouched
    assert worksheet.cell(row=3, column=columns['latitude']).value is None
